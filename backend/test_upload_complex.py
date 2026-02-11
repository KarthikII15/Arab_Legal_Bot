import requests
import io
from docx import Document

BASE_URL = "http://127.0.0.1:5000"

def create_docx():
    doc = Document()
    doc.add_heading('قضية تجارية', 0)
    doc.add_paragraph('الوقائع: تقدم المدعي بدعوى يطالب فيها بمبلغ 5000 ريال.')
    doc.add_paragraph('الأسباب: ثبت للدائرة صحة الدعوى.')
    doc.add_paragraph('الحكم: إلزام المدعى عليه بالمبلغ.')
    
    f = io.BytesIO()
    doc.save(f)
    return f.getvalue()

def test_upload_analyze_docx():
    print("\nTesting /upload-analyze with DOCX...")
    docx_bytes = create_docx()
    files = {'file': ('case.docx', docx_bytes, 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')}
    try:
        r = requests.post(f"{BASE_URL}/upload-analyze", files=files)
        print(f"Status: {r.status_code}")
        if r.status_code == 200:
            print("PASS")
            print(f"Classification: {r.json()['classification']['name_en']}")
        else:
            print(f"FAIL: {r.text}")
    except Exception as e:
        print(f"Error: {e}")

if __name__ == "__main__":
    test_upload_analyze_docx()
